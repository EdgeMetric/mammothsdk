"""Update method, workbook upload results and file settings against a real DEV instance.

No mocks: each test uploads real files, reads the stored state back through the API and
deletes what it created, even when an assertion fails.

Credentials: the ``MAMMOTH_*`` group of ``conftest.py`` (``adv_client``), a token or a
key and secret; the tests skip when it is not set.
"""

from __future__ import annotations

import contextlib
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from mammoth import MammothClient
from mammoth.models.datasets import DataAdditionMode

pytestmark = pytest.mark.asyncio(loop_scope="session")

#: What the backend stores for a dataset with no settings of its own
#: (mvc-service ``SchemaPropertiesDefaultValues``).
_RELEASE_DEFAULTS: dict[str, Any] = {
    "data_addition_mode": "combine",
    "delete_after_append": True,
    "strict_data_type_check": True,
    "convert_columns_to_text": False,
    "strict_column_length_check": False,
    "match_existing_target_schema": True,
    "convert_incompatible_data_to_null": False,
}
_SHEETS = ("Jan", "Feb", "Mar")


@pytest.fixture(scope="module")
def client(adv_client: MammothClient) -> MammothClient:
    """The suite's ``MAMMOTH_*`` client, under the name these tests use."""
    return adv_client


async def _stored_schema_properties(client: MammothClient, dataset_id: int) -> dict[str, Any]:
    """The dataset's ``configurations.schema_properties``, read with a real GET."""
    response = await client.datasets.get(dataset_id, fields="id,configurations")
    dataset = response.get("dataset", response)
    return (dataset.get("configurations") or {}).get("schema_properties") or {}


def _settings_subset(stored: dict[str, Any]) -> dict[str, Any]:
    return {key: stored[key] for key in _RELEASE_DEFAULTS if key in stored}


async def test_set_update_method_changes_only_the_mode(
    client: MammothClient, tmp_path: Path
) -> None:
    """Replace keeps every other stored setting; an empty dataset gets the release defaults."""
    csv_path = tmp_path / "update_method.csv"
    csv_path.write_text("name,salary\nalice,100\nbob,90\n")
    dataset_id = await client.files.upload(str(csv_path))
    assert isinstance(dataset_id, int)
    try:
        read = await client.datasets.get_update_method(dataset_id)
        before = _settings_subset(await _stored_schema_properties(client, dataset_id))
        assert read == {
            k: v for k, v in read.items() if k in {*_RELEASE_DEFAULTS, "schema_mapping"}
        }
        expected = {**_RELEASE_DEFAULTS, **before, "data_addition_mode": "replace"}

        result = await client.datasets.set_update_method(dataset_id, DataAdditionMode.REPLACE)

        assert result["data_addition_mode"] == "replace"
        assert result["dataset_id"] == dataset_id
        after = _settings_subset(await _stored_schema_properties(client, dataset_id))
        assert after == expected
        assert (await client.datasets.get_update_method(dataset_id))["data_addition_mode"] == (
            "replace"
        )

        restored = await client.datasets.set_update_method(dataset_id, DataAdditionMode.COMBINE)
        assert restored["data_addition_mode"] == "combine"
        assert _settings_subset(await _stored_schema_properties(client, dataset_id)) == {
            **expected,
            "data_addition_mode": "combine",
        }
    finally:
        with contextlib.suppress(Exception):
            await client.datasets.delete(dataset_id)


async def test_upload_result_reports_a_workbook_waiting_for_its_sheets(
    client: MammothClient, tmp_path: Path
) -> None:
    """A 3-sheet workbook makes no dataset: it waits for the sheet choice."""
    workbook = Workbook()
    workbook.active.title = _SHEETS[0]
    for title in _SHEETS[1:]:
        workbook.create_sheet(title)
    for sheet in workbook.worksheets:
        sheet.append(["name", "salary"])
        sheet.append(["alice", 100])
    path = tmp_path / "three_sheets_update_method.xlsx"
    workbook.save(path)

    outcome: dict[str, Any] = {}
    try:
        outcome = await client.files.upload_result(str(path))
        assert outcome["dataset_ids"] == []
        assert "files_error" not in outcome
        assert len(outcome["files"]) == 1
        entry = outcome["files"][0]
        assert entry["name"] == path.name
        assert entry["action_needed"] == "sheet_selection_required"
        names = [sheet.get("sheet_name") or sheet.get("name") for sheet in entry["sheets_info"]]
        assert sorted(names) == sorted(_SHEETS)
    finally:
        for entry in outcome.get("files", []):
            with contextlib.suppress(Exception):
                await client.files.delete(entry["file_id"])
        for dataset_id in outcome.get("dataset_ids", []):
            with contextlib.suppress(Exception):
                await client.datasets.delete(dataset_id)


async def test_upload_result_gives_each_same_named_upload_its_own_file_id(
    client: MammothClient, tmp_path: Path
) -> None:
    """The same workbook name uploaded twice within a minute: each reports its own file id."""
    workbook = Workbook()
    workbook.active.title = _SHEETS[0]
    for title in _SHEETS[1:]:
        workbook.create_sheet(title)
    for sheet in workbook.worksheets:
        sheet.append(["name", "salary"])
        sheet.append(["alice", 100])
    path = tmp_path / "same_name_twice_update_method.xlsx"
    workbook.save(path)

    outcomes: list[dict[str, Any]] = []
    try:
        for _ in range(2):
            outcomes.append(await client.files.upload_result(str(path)))
        file_ids = []
        for outcome in outcomes:
            assert "files_error" not in outcome
            assert len(outcome["files"]) == 1
            entry = outcome["files"][0]
            assert entry["name"] == path.name
            assert entry["action_needed"] == "sheet_selection_required"
            # The id the upload job itself named for this upload's file.
            assert [entry["file_id"]] == [
                info["file_object_id"] for info in outcome["nested_job_ids"]
            ]
            file_ids.append(entry["file_id"])
        assert len(set(file_ids)) == 2
    finally:
        for outcome in outcomes:
            for entry in outcome.get("files", []):
                with contextlib.suppress(Exception):
                    await client.files.delete(entry["file_id"])


async def test_file_settings_update_with_one_setting_keeps_the_others(
    client: MammothClient, tmp_path: Path
) -> None:
    """Changing only ``initial_skip_count`` leaves delimiter, header and quote as they were."""
    csv_path = tmp_path / "two_title_lines_update_method.csv"
    csv_path.write_text("Quarterly report\nprepared by finance\nname,salary\nalice,100\nbob,90\n")
    dataset_id = await client.files.upload(str(csv_path))
    assert isinstance(dataset_id, int)
    try:
        before = (await client.datasets.get_file_settings(dataset_id))["info"]
        kept = ("delimiter", "has_header", "quotechar")

        response = await client.datasets.file_settings_update(dataset_id, initial_skip_count=2)
        job_id = response.get("job_id") if isinstance(response, dict) else None
        if job_id:
            await client.jobs.wait_for_job(job_id)

        after = (await client.datasets.get_file_settings(dataset_id))["info"]
        assert after["initial_skip_count"] == 2
        for key in kept:
            assert after[key] == before[key], key
    finally:
        with contextlib.suppress(Exception):
            await client.datasets.delete(dataset_id)


async def _make_dataset(client: MammothClient, tmp_path: Path, name: str, rows: int) -> int:
    """Upload a ``name,salary`` CSV of ``rows`` rows as a new dataset and return its id."""
    path = tmp_path / name
    path.write_text("name,salary\n" + "".join(f"person{i},{100 + i}\n" for i in range(rows)))
    dataset_id = await client.files.upload(str(path))
    assert isinstance(dataset_id, int)
    return dataset_id


def _csv_of(tmp_path: Path, name: str, rows: int) -> Path:
    """Write a ``name,salary`` CSV of ``rows`` rows and return its path."""
    path = tmp_path / name
    path.write_text("name,salary\n" + "".join(f"extra{i},{200 + i}\n" for i in range(rows)))
    return path


async def test_append_under_add_reports_its_own_batches_and_rows(
    client: MammothClient, tmp_path: Path
) -> None:
    """An upload under ``add`` reports the batch it made and the rows it added."""
    dataset_id = await _make_dataset(client, tmp_path, "append_add_base.csv", rows=2)
    try:
        extra = _csv_of(tmp_path, "append_add_more.csv", rows=3)

        report = await client.appends.append(str(extra), dataset_id, mode="add")

        assert report["appended_to"] == dataset_id
        assert report["update_method"]["mode"] == "add"
        assert "append_failed" not in report
        assert "append_pending" not in report
        assert report["rows_before"] == 2
        assert report["rows_added"] == 3
        listed = {row["id"] for row in await client.appends.list_batches(dataset_id)}
        assert report["batch_ids"]
        assert set(report["batch_ids"]) <= listed
    finally:
        with contextlib.suppress(Exception):
            await client.datasets.delete(dataset_id)


async def test_append_under_replace_sets_the_mode_and_reports_it(
    client: MammothClient, tmp_path: Path
) -> None:
    """An upload under ``replace`` leaves the dataset's update method as replace."""
    dataset_id = await _make_dataset(client, tmp_path, "append_replace_base.csv", rows=3)
    try:
        extra = _csv_of(tmp_path, "append_replace_new.csv", rows=1)

        report = await client.appends.append(str(extra), dataset_id, mode="replace")

        assert report["update_method"]["mode"] == "replace"
        assert report["update_method"]["previous_mode"] == "add"
        assert "append_failed" not in report
        assert report["rows_after"] == 1
        stored = await client.datasets.get_update_method(dataset_id)
        assert stored["data_addition_mode"] == "replace"
    finally:
        with contextlib.suppress(Exception):
            await client.datasets.delete(dataset_id)


async def test_restore_mode_puts_back_the_mode_a_prepare_changed(
    client: MammothClient, tmp_path: Path
) -> None:
    """``restore_mode`` undoes a change ``prepare`` made, and does nothing when none was."""
    dataset_id = await _make_dataset(client, tmp_path, "restore_mode_base.csv", rows=2)
    try:
        unchanged = await client.appends.prepare(dataset_id, "add")
        assert unchanged["mode_changed"] is False
        assert await client.appends.restore_mode(unchanged) is None

        plan = await client.appends.prepare(dataset_id, "replace")
        assert plan["mode_changed"] is True
        assert plan["update_method"]["previous_mode"] == "add"
        assert (await client.datasets.get_update_method(dataset_id))["data_addition_mode"] == (
            "replace"
        )

        assert await client.appends.restore_mode(plan) is None

        assert (await client.datasets.get_update_method(dataset_id))["data_addition_mode"] == (
            "combine"
        )
    finally:
        with contextlib.suppress(Exception):
            await client.datasets.delete(dataset_id)
